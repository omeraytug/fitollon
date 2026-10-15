"""Parse MacroFactor's in-app data export (More → Data Management → Data Export), an .xlsx workbook.

Each sheet has a Date column followed by value columns such as "Protein (g)". Every filled cell becomes a
daily metric with provider "macrofactor". Columns that also exist in Apple Health get the Apple Health
metric name so both sources line up; the rest get a snake_case name ("Vitamin D (mcg)" -> "vitamin_d").
"""

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from fitollon.ingest.health_auto_export import MetricRow

PROVIDER = "macrofactor"

# Goal ranges (floor/target/ceiling per nutrient), not daily values.
SKIP_SHEETS = {"Micronutrient Goals"}

# Column name -> metric name, for quantities Apple Health also has, plus MacroFactor's own calculations.
METRIC_NAMES = {
    "Calories": "dietary_energy",
    "Fat": "total_fat",
    "Carbs": "carbohydrates",
    "Protein": "protein",
    "Fiber": "fiber",
    "Weight": "weight_body_mass",
    "Fat Percent": "body_fat_percentage",
    "Steps": "step_count",
    "Trend Weight": "weight_trend",
    "Expenditure": "expenditure",
    "Visual Body Fat Assessment": "visual_body_fat",
}

# Units for columns whose header doesn't name one.
DEFAULT_UNITS = {"Steps": "count", "Expenditure": "kcal", "Fat Percent": "%", "Visual Body Fat Assessment": "%"}

HEADER = re.compile(r"^(?P<name>.*?)\s*(?:\((?P<unit>[^)]*)\))?\s*$")


def _metric(header: str) -> tuple[str, str | None]:
    match = HEADER.match(header.strip())
    name, unit = match["name"], match["unit"]
    metric = METRIC_NAMES.get(name) or re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
    return metric, unit or DEFAULT_UNITS.get(name)


def _day(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return None
    return None


def parse_workbook(path: Path) -> list[MetricRow]:
    rows: dict[tuple[date, str], MetricRow] = {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            if sheet.title in SKIP_SHEETS:
                continue
            values = sheet.iter_rows(values_only=True)
            header = next(values, None)
            if not header or str(header[0]).strip().lower() != "date":
                continue
            columns = [_metric(str(h)) if h else None for h in header[1:]]
            for row in values:
                day = _day(row[0]) if row else None
                if day is None:
                    continue
                for column, cell in zip(columns, row[1:], strict=False):
                    if column is None or isinstance(cell, bool) or not isinstance(cell, int | float):
                        continue
                    metric, unit = column
                    rows[(day, metric)] = MetricRow(day=day, metric=metric, value=float(cell), unit=unit)
    finally:
        workbook.close()
    return list(rows.values())
